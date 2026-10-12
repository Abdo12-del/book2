#!/usr/bin/env python3
"""مُصدّر DOCX: يبني ملف Word قابل للتحرير بالكامل بأنماط حقيقية.

المزايا:
  • أقسام حقيقية: غلاف ← تمهيد (ترقيم روماني) ← متن (ترقيم يبدأ من 1).
  • أنماط Word قياسية (Heading 1-4 / Body Text / Caption / Quote ...) مضبوطة عربياً.
  • اتجاه RTL على مستوى القسم (w:bidi) والفقرات، وهوامش متقابلة (Mirror Margins).
  • فهرس محتويات تلقائي بحقل TOC، وترقيم صفحات بحقل PAGE، وترويسة بحقل STYLEREF.
  • جداول وصور وصناديق ملاحظات مع تعليقات مُرقّمة.
"""
from __future__ import annotations

import os

import docx
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import arabic as A
from manuscript import Block, _inline
from theme import MM, load as load_theme

# ترتيب العناصر داخل w:pPr / w:rPr / w:sectPr وفق مخطط OOXML (مهم لصلاحية الملف)
PPR_AFTER_BIDI = ("w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                  "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                  "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
                  "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
RPR_AFTER_RTL = ("w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")
RPR_AFTER_LANG = ("w:eastAsianLayout", "w:specVanish", "w:oMath")
RPR_AFTER_RFONTS = ("w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike",
                    "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint", "w:noProof",
                    "w:snapToGrid", "w:vanish", "w:webHidden", "w:color", "w:spacing", "w:w",
                    "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect",
                    "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em",
                    "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")
SECTPR_AFTER_BIDI = ("w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange")
PPR_AFTER_SHD = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                 "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                 "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                 "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                 "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
                 "w:sectPr", "w:pPrChange")
RPR_AFTER_BCS = ("w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike", "w:dstrike", "w:outline",
                 "w:shadow", "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish",
                 "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz",
                 "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText",
                 "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
                 "w:specVanish", "w:oMath")
RPR_AFTER_SZCS = ("w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign",
                  "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")
SECTPR_AFTER_PGNUM = ("w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
                      "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
                      "w:printerSettings", "w:sectPrChange")

MM_TO_CM = 0.1


def _sub(parent, tag: str, successors=(), **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    parent.insert_element_before(el, *successors)
    return el


def make_rtl(p, rtl: bool = True) -> None:
    """ضبط اتجاه الفقرة (w:bidi) — الطريقة الصحيحة لجعل الفقرة عربية الاتجاه."""
    ppr = p._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        _sub(ppr, "w:bidi", PPR_AFTER_BIDI, val="1" if rtl else "0")
    else:
        ppr.find(qn("w:bidi")).set(qn("w:val"), "1" if rtl else "0")
    if rtl:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def style_run(run, font_name: str, size_pt: float | None = None, rtl: bool = True,
              bold: bool | None = None, italic: bool | None = None,
              color: str | None = None) -> None:
    rpr = run._r.get_or_add_rPr()
    if rpr.find(qn("w:rFonts")) is None:
        _sub(rpr, "w:rFonts", RPR_AFTER_RFONTS)
    rf = rpr.find(qn("w:rFonts"))
    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
        rf.set(qn(f"w:{a}"), font_name)
    if bold is not None:
        run.bold = bold
        b = rpr.find(qn("w:bCs"))
        if b is None:
            _sub(rpr, "w:bCs", RPR_AFTER_BCS, val="1" if bold else "0")
        else:
            b.set(qn("w:val"), "1" if bold else "0")
    if italic is not None:
        run.italic = italic
    if size_pt is not None:
        run.font.size = Pt(size_pt)
        for tag in ("w:szCs",):
            el = rpr.find(qn(tag))
            if el is None:
                _sub(rpr, tag, RPR_AFTER_SZCS)
            rpr.find(qn(tag)).set(qn("w:val"), str(int(size_pt * 2)))
    if color:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    if rtl and rpr.find(qn("w:rtl")) is None:
        _sub(rpr, "w:rtl", RPR_AFTER_RTL, val="1")
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = _sub(rpr, "w:lang", RPR_AFTER_LANG)
    lang.set(qn("w:val"), "ar-SA")
    lang.set(qn("w:bidi"), "ar-SA")


def add_field(p, instr: str, placeholder: str = "", font_name: str = "Amiri",
              size_pt: float = 10, rtl: bool = True) -> None:
    """إدراج حقل Word حقيقي (PAGE / TOC / STYLEREF / NUMPAGES)."""
    def run():
        return p.add_run()

    r = run()
    style_run(r, font_name, size_pt, rtl)
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    fc.set(qn("w:dirty"), "true")
    r._r.append(fc)

    r = run()
    style_run(r, font_name, size_pt, rtl)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)

    r = run()
    style_run(r, font_name, size_pt, rtl)
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "separate")
    r._r.append(fc)

    r = run()
    style_run(r, font_name, size_pt, rtl)
    r.text = placeholder

    r = run()
    style_run(r, font_name, size_pt, rtl)
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    r._r.append(fc)


def set_cell_bg(cell, hex_color: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tcpr.append(shd)


def set_cell_border(cell, side: str, size: int = 18, color: str = "8C6D3F") -> None:
    """إطار خلية بسماكة مُعطاة (الحجم بوحدة 1/8 نقطة)."""
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.insert_element_before(borders, "w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
                                   "w:tcFitText", "w:vAlign", "w:hideMark")
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:color"), color)
    borders.append(el)


def para_spacing(p, before=0, after=0, line=None, line_rule=None) -> None:
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
        if line_rule:
            pf.line_spacing_rule = line_rule


def keep_with_next(p, on: bool = True) -> None:
    ppr = p._p.get_or_add_pPr()
    if ppr.find(qn("w:keepNext")) is None:
        _sub(ppr, "w:keepNext", ("w:keepLines", "w:pageBreakBefore", "w:framePr",
                                 "w:widowControl", "w:numPr", "w:suppressLineNumbers",
                                 "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens",
                                 "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
                                 "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi"))


def setup_section(sec, theme: dict, rtl: bool = True) -> None:
    p = theme["page"]
    sec.page_width = Cm(p["width_pt"] / MM / 10)
    sec.page_height = Cm(p["height_pt"] / MM / 10)
    inner = (p["margin_outer_mm"] + p["gutter_mm"]) * MM_TO_CM
    outer = p["margin_outer_mm"] * MM_TO_CM
    # مع تفعيل "الهوامش المتقابلة" يصبح اليسار = الداخلي، واليمين = الخارجي
    sec.left_margin = Cm(inner)
    sec.right_margin = Cm(outer)
    sec.top_margin = Cm(p["margin_top_mm"] * MM_TO_CM)
    sec.bottom_margin = Cm(p["margin_bottom_mm"] * MM_TO_CM)
    sec.header_distance = Cm(p["header_from_top_mm"] * MM_TO_CM)
    sec.footer_distance = Cm(p["footer_from_bottom_mm"] * MM_TO_CM)
    sect_pr = sec._sectPr
    if sect_pr.find(qn("w:bidi")) is None:
        _sub(sect_pr, "w:bidi", SECTPR_AFTER_BIDI, val="1" if rtl else "0")
    if rtl and sect_pr.find(qn("w:rtlGutter")) is None:
        _sub(sect_pr, "w:rtlGutter", ("w:docGrid", "w:printerSettings", "w:sectPrChange"), val="1")


def set_page_numbering(sec, fmt: str = "decimal", start: int = 1) -> None:
    sect_pr = sec._sectPr
    el = sect_pr.find(qn("w:pgNumType"))
    if el is None:
        el = _sub(sect_pr, "w:pgNumType", SECTPR_AFTER_PGNUM)
    el.set(qn("w:fmt"), fmt)
    el.set(qn("w:start"), str(start))


def _shade(paragraph, hex_color: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    ppr.insert_element_before(shd, *PPR_AFTER_SHD)


# --------------------------------------------------------------------------- الأنماط
def setup_styles(doc, theme: dict, fonts: dict[str, str]) -> None:
    s, c = theme["sizes"], theme["colors"]
    rtl = theme["direction"] == "rtl"
    align = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    body_font, head_font = fonts["body"], fonts["heading"]

    def S(name, base=None):
        try:
            st = doc.styles[name]
        except KeyError:
            st = doc.styles.add_style(name, 1 if base in (None,) else 1)
        return st

    # النمط الافتراضي
    normal = doc.styles["Normal"]
    normal.font.name = body_font
    normal.font.size = Pt(s["body_pt"])
    rpr = normal.element.get_or_add_rPr()
    if rpr.find(qn("w:rFonts")) is None:
        _sub(rpr, "w:rFonts", RPR_AFTER_RFONTS)
    rf = rpr.find(qn("w:rFonts"))
    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
        rf.set(qn(f"w:{a}"), body_font)
    if rtl and rpr.find(qn("w:rtl")) is None:
        _sub(rpr, "w:rtl", RPR_AFTER_RTL, val="1")
    normal.paragraph_format.line_spacing = s["body_leading_pt"] / s["body_pt"]
    normal.paragraph_format.widow_control = True      # منع السطر اليتيم/الأسطر المعلّقة
    normal.paragraph_format.keep_with_next = False
    normal.paragraph_format.space_after = Pt(s["body_leading_pt"] * 0.35)
    normal.paragraph_format.alignment = align

    heads = [("Heading 1", s["h1_pt"], 18, 10), ("Heading 2", s["h2_pt"], 14, 6),
             ("Heading 3", s["h3_pt"], 11, 5), ("Heading 4", s["h4_pt"], 9, 4)]
    for name, size, before, after in heads:
        st = doc.styles[name]
        st.font.name = head_font
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(c["heading"].lstrip("#"))
        rpr = st.element.get_or_add_rPr()
        if rpr.find(qn("w:rFonts")) is None:
            _sub(rpr, "w:rFonts", RPR_AFTER_RFONTS)
        rf = rpr.find(qn("w:rFonts"))
        for a in ("ascii", "hAnsi", "cs", "eastAsia"):
            rf.set(qn(f"w:{a}"), head_font)
        pf = st.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if name == "Heading 1" else align
        pf.keep_with_next = True
        pf.widow_control = True

    cap = doc.styles["Caption"]
    cap.font.name = body_font
    cap.font.size = Pt(s["caption_pt"])
    cap.font.color.rgb = RGBColor.from_string(c["muted"].lstrip("#"))
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(10)

    q = doc.styles["Quote"]
    q.font.name = fonts.get("body_italic", body_font)
    q.font.size = Pt(s["body_pt"] * 0.95)
    q.font.italic = True
    q.font.color.rgb = RGBColor.from_string(c["muted"].lstrip("#"))
    q.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for nm in ("List Bullet", "List Number", "List Paragraph", "Body Text"):
        try:
            st = doc.styles[nm]
            st.font.name = body_font
            st.font.size = Pt(s["body_pt"])
            st.paragraph_format.alignment = align
        except KeyError:
            pass

    # أنماط مخصصة للكتاب
    def custom(name: str, font: str, size: float, color: str,
               align_: WD_ALIGN_PARAGRAPH, before: float = 0, after: float = 6,
               bold: bool = False, line: float | None = None):
        try:
            st = doc.styles[name]
        except KeyError:
            st = doc.styles.add_style(name, 1)  # 1 = paragraph style
        st.font.name = font
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        rpr = st.element.get_or_add_rPr()
        if rpr.find(qn("w:rFonts")) is None:
            _sub(rpr, "w:rFonts", RPR_AFTER_RFONTS)
        rf = rpr.find(qn("w:rFonts"))
        for a in ("ascii", "hAnsi", "cs", "eastAsia"):
            rf.set(qn(f"w:{a}"), font)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.alignment = align_
        if line:
            st.paragraph_format.line_spacing = line
        return st

    custom("Book Chapter Label", head_font, s["h1_pt"] * 0.62, c["accent"],
           WD_ALIGN_PARAGRAPH.CENTER, 0, 6)
    custom("Book Callout", body_font, s["body_pt"] * 0.96, c["text"], align)
    custom("Book Callout Title", head_font, s["body_pt"] * 0.95, c["accent"], align, 0, 3, True)
    custom("Book Code", fonts["mono"], s["body_pt"] * 0.82, c["text"],
           WD_ALIGN_PARAGRAPH.LEFT, 6, 8)
    custom("Book TOC", body_font, s["toc_pt"], c["text"], align, 0, 2)


# --------------------------------------------------------------------------- الكُتل
class DocxBuilder:
    def __init__(self, doc, theme, fonts):
        self.doc, self.th, self.f = doc, theme, fonts
        self.chapter = 0
        self.table_n = 0
        self.fig_n = 0
        self.rtl = theme["direction"] == "rtl"
        self.align = WD_ALIGN_PARAGRAPH.RIGHT if self.rtl else WD_ALIGN_PARAGRAPH.LEFT
        self.sizes, self.colors = theme["sizes"], theme["colors"]
        self.text_width_cm = (theme["page"]["width_pt"] * MM / 10
                              - (theme["page"]["margin_outer_mm"] * 2 + theme["page"]["gutter_mm"]) * MM_TO_CM)

    # فقرة بتنسيق مضمّن
    def rich(self, text: str, style_name: str | None = None, font=None, size=None,
             color=None, align=None, rtl=None, before=None, after=None):
        p = self.doc.add_paragraph(style=style_name)
        make_rtl(p, self.rtl if rtl is None else rtl)
        if align is not None:
            p.alignment = align
        if before is not None or after is not None:
            para_spacing(p, before or 0, after or 0)
        for kind, seg in _inline(A.normalize_numerals(text)):
            r = p.add_run()
            style_run(r, font or self.f["body"], size or self.sizes["body_pt"],
                      self.rtl, bold=(kind == "b"), italic=(kind == "i"),
                      color=color or self.colors["text"])
            if kind == "c":
                style_run(r, self.f["mono"], (size or self.sizes["body_pt"]) * 0.85, False)
            r.text = seg
        return p

    def heading(self, b: Block):
        if b.level == 1:
            self.chapter += 1
            self.table_n = self.fig_n = 0
            # الافتتاحية: سطر الرقم المولَّد (اختياري) ثم عنوان الفصل
            pre = self.th.get("meta", {}).get("chapter_label", None) or "الفصل"
            drop = bool(self.th["layout"].get("drop_chapter_number"))
            if not drop:
                p = self.doc.add_paragraph()
                make_rtl(p, self.rtl)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(self.th["layout"]["chapter_head_space_before_pt"])
                keep_with_next(p)
                r = p.add_run()
                style_run(r, self.f["heading"], self.sizes["h1_pt"] * 0.62, self.rtl,
                          color=self.colors["accent"])
                r.text = A.normalize_numerals(f"{pre} " + _chapter_word(self.chapter))
            p2 = self.doc.add_paragraph(style="Heading 1")
            make_rtl(p2, self.rtl)
            if drop:      # بلا سطر رقم: نُبقي هبوط العنوان من أعلى الصفحة
                p2.paragraph_format.space_before = Pt(
                    self.th["layout"]["chapter_head_space_before_pt"])
            r = p2.add_run()
            style_run(r, self.f["heading"], self.sizes["h1_pt"], self.rtl,
                      bold=True, color=self.colors["heading"])
            r.text = A.normalize_numerals(b.text)
            keep_with_next(p2)
            return
        style = f"Heading {min(b.level, 4)}"
        p = self.doc.add_paragraph(style=style)
        make_rtl(p, self.rtl)
        r = p.add_run()
        style_run(r, self.f["heading"], self.sizes.get(f"h{min(b.level,4)}_pt",
                                                       self.sizes["body_pt"]),
                  self.rtl, bold=True, color=self.colors["heading"])
        r.text = A.normalize_numerals(b.text)
        keep_with_next(p)

    def paragraph(self, b: Block):
        self.rich(b.text, align=self.align,
                  after=self.sizes["body_leading_pt"] * 0.35)

    def bullet_list(self, b: Block):
        ordered = b.data.get("ordered")
        counter = 1
        for it in b.data["items"]:
            lvl = min(it["level"], 2)
            st = ("List Number" if ordered else "List Bullet") + (f" {lvl + 1}" if lvl else "")
            try:
                p = self.doc.add_paragraph(style=st)
                _ = self.doc.styles[st]
            except KeyError:
                p = self.doc.add_paragraph()
            make_rtl(p, self.rtl)
            p.paragraph_format.space_after = Pt(3)
            if ordered:
                marker = A.format_number(counter, A.numeral_style()) + ".\u00a0"
                counter += 1
            else:
                marker = ("•\u00a0" if lvl == 0 else "–\u00a0")
            r = p.add_run()
            style_run(r, self.f["body"], self.sizes["body_pt"], self.rtl)
            r.text = marker
            for kind, seg in _inline(it["text"]):
                r = p.add_run()
                style_run(r, self.f["body"], self.sizes["body_pt"], self.rtl,
                          bold=(kind == "b"), italic=(kind == "i"))
                r.text = seg
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def table_block(self, b: Block):
        header, rows = b.data["header"], b.data["rows"]
        ncol = max([len(header)] + [len(r) for r in rows])
        header = header + [""] * (ncol - len(header))
        rows = [r + [""] * (ncol - len(r)) for r in rows]
        t = self.doc.add_table(rows=len(rows) + 1, cols=ncol)
        try:
            t.style = self.doc.styles["Table Grid"]
        except KeyError:
            pass
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblpr = t._tbl.tblPr
        if self.rtl and tblpr.find(qn("w:bidiVisual")) is None:
            bv = OxmlElement("w:bidiVisual")
            tblpr.append(bv)
        widths = [int(self.text_width_cm / ncol * 567) for _ in range(ncol)]
        for row in t.rows:
            for i, cell in enumerate(row.cells):
                cell.width = Cm(self.text_width_cm / ncol)
        for j, txt in enumerate(header):
            cell = t.rows[0].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            make_rtl(p, self.rtl)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for kind, seg in _inline(A.normalize_numerals(txt)):
                r = p.add_run()
                style_run(r, self.f["heading"] if kind != "c" else self.f["mono"],
                          self.sizes["table_pt"] * (0.92 if kind == "c" else 1.0),
                          self.rtl, bold=True, color=self.colors["heading"])
                r.text = seg
            set_cell_bg(cell, self.colors["table_header_bg"])
        for i, row in enumerate(rows):
            for j, txt in enumerate(row):
                cell = t.rows[i + 1].cells[j]
                cell.text = ""
                p = cell.paragraphs[0]
                make_rtl(p, self.rtl)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for kind, seg in _inline(A.normalize_numerals(txt)):
                    r = p.add_run()
                    style_run(r, self.f["body"] if kind != "c" else self.f["mono"],
                              self.sizes["table_pt"] * (0.92 if kind == "c" else 1.0),
                              self.rtl, bold=(kind == "b"), italic=(kind == "i"))
                    r.text = seg
                if self.th["objects"]["table_style"] == "zebra" and i % 2 == 1:
                    set_cell_bg(cell, self.colors["table_zebra_bg"])
        self.doc.add_paragraph()
        self._caption(self.th["objects"]["caption_prefix_table"], b.data.get("caption", ""), "table")

    def figure(self, b: Block):
        src = b.data.get("src", "")
        if not os.path.exists(src):
            self.rich(f"[صورة غير موجودة: {src}]", align=WD_ALIGN_PARAGRAPH.CENTER)
            return
        max_w = Cm(self.text_width_cm * self.th["objects"]["image_max_width_ratio"])
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(src, width=max_w)
        self._caption(self.th["objects"]["caption_prefix_figure"],
                      b.data.get("caption") or b.text or "", "figure")

    def _caption(self, prefix: str, text: str, kind: str):
        num_style = A.numeral_style()
        if kind == "table":
            self.table_n += 1
            idx = self.table_n
        else:
            self.fig_n += 1
            idx = self.fig_n
        num = (f'{A.format_number(self.chapter, num_style)}-{A.format_number(idx, num_style)}'
               if self.th["objects"]["caption_numbering"] == "chapter"
               else A.format_number(idx, num_style))
        p = self.doc.add_paragraph(style="Caption")
        make_rtl(p, self.rtl)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        style_run(r, self.f["body"], self.sizes["caption_pt"], self.rtl,
                  color=self.colors["muted"])
        r.text = f"{prefix} ({num}) {A.normalize_numerals(text)}".strip()

    def callout(self, b: Block):
        titles = {"note": "ملاحظة", "tip": "نصيحة", "warning": "تنبيه", "example": "مثال",
                  "quote": "اقتباس", "exercise": "تمرين", "definition": "تعريف", "key": "مفتاح"}
        kind = b.data.get("kind", "note")
        title = b.data.get("title") or titles.get(kind, "ملاحظة")
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblpr = t._tbl.tblPr
        if self.rtl and tblpr.find(qn("w:bidiVisual")) is None:
            tblpr.append(OxmlElement("w:bidiVisual"))
        cell = t.rows[0].cells[0]
        cell.width = Cm(self.text_width_cm)
        set_cell_bg(cell, self.colors["callout_bg"])
        set_cell_border(cell, "right" if self.rtl else "left", 18,
                        self.colors["callout_border"].lstrip("#"))
        first = True
        for sub in (b.data.get("blocks") or [Block("p", text=b.text)]):
            if sub.t == "p":
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                make_rtl(p, self.rtl)
                for kind_, seg in _inline(sub.text):
                    r = p.add_run()
                    style_run(r, self.f["body"], self.sizes["body_pt"] * 0.96, self.rtl,
                              bold=(kind_ == "b"), italic=(kind_ == "i"))
                    r.text = seg
        self.doc.add_paragraph()

    def code(self, b: Block):
        for line in b.text.split("\n"):
            p = self.doc.add_paragraph(style="Book Code")
            make_rtl(p, False)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run()
            style_run(r, self.f["mono"], self.sizes["body_pt"] * 0.82, False)
            r.text = A.normalize_numerals(line) or " "
            _shade(p, "F2F2F2")

    def build(self, blocks):
        for b in blocks:
            if b.t == "h":
                if b.level == 1:
                    p = self.doc.add_paragraph()
                    p.add_run().add_break(WD_BREAK.PAGE)
                self.heading(b)
            elif b.t == "p":
                self.paragraph(b)
            elif b.t == "list":
                self.bullet_list(b)
            elif b.t == "table":
                self.table_block(b)
            elif b.t == "figure":
                self.figure(b)
            elif b.t == "callout":
                self.callout(b)
            elif b.t == "code":
                self.code(b)
            elif b.t == "rule":
                p = self.doc.add_paragraph()
                make_rtl(p, self.rtl)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run()
                style_run(r, self.f["body"], 10, self.rtl, color=self.colors["accent"])
                r.text = "◆ ──────── ◆"
                para_spacing(p, 8, 10)
            elif b.t == "pagebreak":
                self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _chapter_word(n: int) -> str:
    from build_pdf import chapter_word
    return chapter_word(n)


def build_docx(blocks, meta: dict, out_path: str, theme_path: str | None = None,
               cover: bool = True, toc: bool = True) -> str:
    from build_pdf import register_fonts
    theme = load_theme(theme_path)
    A.set_numeral_style(theme.get("numerals", "western"))
    A.set_base_dir(theme.get("direction", "rtl"))   # اتجاه الفقرة في Bidi (R للعربية)
    if theme["running"].get("front_matter_numbering") != "roman":
        theme["running"]["front_matter_numbering"] = A.numeral_style()
    fonts = register_fonts(theme)
    rtl = theme["direction"] == "rtl"

    doc = Document()
    cp = doc.core_properties
    cp.title = meta.get("title", "")
    cp.author = meta.get("author", "")
    cp.subject = meta.get("subtitle", "")
    cp.comments = meta.get("description", "")

    setup_styles(doc, theme, fonts)
    sec = doc.sections[0]
    setup_section(sec, theme, rtl)

    # هوامش متقابلة + تحديث الحقول تلقائياً عند الفتح
    try:
        doc.settings.even_and_odd_headers = True
    except Exception:  # noqa: BLE001
        pass
    settings_el = doc.settings.element
    if settings_el.find(qn("w:mirrorMargins")) is None:
        _sub(settings_el, "w:mirrorMargins",
             ("w:alignBordersAndEdges", "w:bordersDoNotSurroundHeader",
              "w:bordersDoNotSurroundFooter", "w:gutterAtTop", "w:hideSpellingErrors"),
             val="1")
    if settings_el.find(qn("w:updateFields")) is None:
        _sub(settings_el, "w:updateFields",
             ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars",
              "w:rsids", "w:mathPr", "w:themeFontLang", "w:clrSchemeMapping",
              "w:shapeDefaults", "w:decimalSymbol", "w:listSeparator"), val="true")

    s = theme["sizes"]

    # ---------- الغلاف
    if cover:
        for f, txt in (("series", "cover_subtitle"), ("title", "cover_title")):
            pass
        if meta.get("series"):
            p = doc.add_paragraph()
            make_rtl(p, rtl)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run()
            style_run(r, fonts["body"], 15, rtl, color=theme["colors"]["muted"])
            r.text = meta["series"]
        p = doc.add_paragraph()
        make_rtl(p, rtl)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para_spacing(p, 60, 10)
        r = p.add_run()
        style_run(r, fonts["heading"], 28, rtl, bold=True, color=theme["colors"]["heading"])
        r.text = meta.get("title", "")
        if meta.get("subtitle"):
            p = doc.add_paragraph()
            make_rtl(p, rtl)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run()
            style_run(r, fonts["body"], 15, rtl, color=theme["colors"]["muted"])
            r.text = meta["subtitle"]
        p = doc.add_paragraph()
        make_rtl(p, rtl)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # لوحة الغلاف (اختيارية) في الوورد أيضًا
        _cv = meta.get("cover_image", "")
        if _cv:
            _cvp = _cv if os.path.isabs(_cv) else os.path.join(
                os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "book", _cv)
            if os.path.exists(_cvp):
                p = doc.add_paragraph()
                make_rtl(p, rtl)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run()
                r.add_picture(_cvp, width=Cm(5.5))
        para_spacing(p, 50, 6)
        r = p.add_run()
        style_run(r, fonts["heading"], 14, rtl, bold=True, color=theme["colors"]["accent"])
        r.text = meta.get("author", "")
        if meta.get("publisher") or meta.get("year"):
            p = doc.add_paragraph()
            make_rtl(p, rtl)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run()
            style_run(r, fonts["body"], 13, rtl, color=theme["colors"]["muted"])
            r.text = "  •  ".join(x for x in [meta.get("publisher", ""),
                                              str(meta.get("year", ""))] if x)
        # إزالة الترويسة/التذييل عن قسم الغلاف
        for part in (sec.header, sec.footer):
            for p in part.paragraphs:
                p._element.getparent().remove(p._element)

    # ---------- قسم التمهيد
    sec_front = doc.add_section(WD_SECTION.NEW_PAGE)
    setup_section(sec_front, theme, rtl)
    set_page_numbering(sec_front, "lowerRoman", 1)
    for part in (sec_front.header, sec_front.footer,
                 sec_front.even_page_header, sec_front.even_page_footer):
        part.is_linked_to_previous = False
        for p in part.paragraphs:
            p._element.getparent().remove(p._element)
    fp = sec_front.footer.paragraphs[0] if sec_front.footer.paragraphs else sec_front.footer.add_paragraph()
    make_rtl(fp, rtl)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, "PAGE", A.format_number(1, A.numeral_style()), fonts["heading"],
              s["footer_pt"], rtl)

    if meta.get("dedication"):
        p = doc.add_paragraph(style="Quote")
        make_rtl(p, rtl)
        r = p.add_run()
        style_run(r, fonts.get("body_italic", fonts["body"]), s["body_pt"] * 0.95, rtl,
                  italic=True, color=theme["colors"]["muted"])
        r.text = meta["dedication"]
    if meta.get("epigraph"):
        p = doc.add_paragraph(style="Quote")
        make_rtl(p, rtl)
        r = p.add_run()
        style_run(r, fonts.get("body_italic", fonts["body"]), s["body_pt"] * 0.95, rtl,
                  italic=True, color=theme["colors"]["muted"])
        r.text = meta["epigraph"]

    if toc and theme["toc"]["enabled"]:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        depth = theme["toc"]["depth"]
        p = doc.add_paragraph(style="Heading 1")
        make_rtl(p, rtl)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        style_run(r, fonts["heading"], s["h1_pt"] * 0.9, rtl, bold=True,
                  color=theme["colors"]["heading"])
        r.text = "المحتويات"
        p = doc.add_paragraph()
        make_rtl(p, rtl)
        add_field(p, f'TOC \\o "1-{depth}" \\h \\z \\u',
                  "حدّث الفهرس من Word: References ← Update Table",
                  fonts["body"], s["toc_pt"], rtl)

    # ---------- قسم المتن
    sec_body = doc.add_section(WD_SECTION.NEW_PAGE)
    setup_section(sec_body, theme, rtl)
    set_page_numbering(sec_body, "decimal", 1)
    for part in (sec_body.header, sec_body.footer,
                 sec_body.even_page_header, sec_body.even_page_footer):
        part.is_linked_to_previous = False
        for p in list(part.paragraphs):
            p._element.getparent().remove(p._element)

    # ترويسة الصفحات الفردية: عنوان الفصل عبر STYLEREF
    hp = sec_body.header.add_paragraph()
    make_rtl(hp, rtl)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    add_field(hp, 'STYLEREF "Heading 1"', "", fonts["body"], s["header_pt"], rtl)
    hp2 = sec_body.even_page_header.add_paragraph()
    make_rtl(hp2, rtl)
    hp2.alignment = WD_ALIGN_PARAGRAPH.LEFT if rtl else WD_ALIGN_PARAGRAPH.RIGHT
    r = hp2.add_run()
    style_run(r, fonts["body"], s["header_pt"], rtl, color=theme["colors"]["muted"])
    r.text = meta.get("title", "")

    for foot in (sec_body.footer, sec_body.even_page_footer):
        fp = foot.add_paragraph()
        make_rtl(fp, rtl)
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(fp, "PAGE", A.format_number(1, A.numeral_style()), fonts["heading"],
              s["footer_pt"], rtl)

    DocxBuilder(doc, theme, fonts).build(blocks)
    doc.save(out_path)
    return out_path
